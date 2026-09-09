"""
Generador Excel profesional de TAN-Telecom.

Hojas:
    1. Verificaciones de Diseño
    2. Lista de Materiales (BOM)
    3. Parámetros Técnicos
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import ProjectResult

# Colores corporativos (azul industrial / acento ámbar)
FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
FILL_SUB = PatternFill("solid", fgColor="2E75B6")
FILL_ALT = PatternFill("solid", fgColor="D6EAF8")
FILL_WARN = PatternFill("solid", fgColor="FCE4D6")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FONT_HEADER = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
FONT_TITLE = Font(name="Calibri", bold=True, size=14, color="1F4E79")
FONT_NORMAL = Font(name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _autosize(ws, min_width: int = 10, max_width: int = 55) -> None:
    """Ajusta anchos de columna de forma heurística."""
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            if cell.value is None:
                continue
            length = max(length, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, length + 2))


def _style_header_row(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN


def _write_verifications(ws, result: ProjectResult) -> None:
    ws["A1"] = "TAN-Telecom — Verificaciones de Diseño (IEC 61439-1 / Anexo D)"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:H1")
    ws["A2"] = (
        f"Proyecto: {result.project_id} | Envolvente: {result.enclosure.codigo} | "
        f"Forma: {result.separation.forma.value} | Familia: {result.separation.familia_tan}"
    )
    ws["A2"].font = FONT_NORMAL
    ws.merge_cells("A2:H2")

    headers = [
        "N°",
        "Características a verificar",
        "Artículo",
        "Ensayos",
        "Comparación",
        "Evaluación",
        "Estado de cumplimiento",
        "Observaciones técnicas",
    ]
    start = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=start, column=i, value=h)
    _style_header_row(ws, start, len(headers))

    for idx, item in enumerate(result.verifications):
        r = start + 1 + idx
        values = [
            item.numero,
            item.caracteristica,
            item.articulo,
            item.ensayos,
            item.comparacion,
            item.evaluacion,
            item.estado,
            item.observaciones,
        ]
        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = FONT_NORMAL
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if idx % 2 == 1:
                cell.fill = FILL_ALT
            if c == 7:
                estado_u = str(val).upper()
                if "REVISAR" in estado_u:
                    cell.fill = FILL_WARN
                elif "CUMPLE" in estado_u or "CUBIERTO" in estado_u:
                    cell.fill = FILL_OK

    ws.row_dimensions[start].height = 30
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{start}:H{start + len(result.verifications)}"
    _autosize(ws)
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["H"].width = 45


def _write_bom(ws, result: ProjectResult) -> None:
    ws["A1"] = "TAN-Telecom — Lista de Materiales (BOM tipificada)"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:F1")
    ws["A2"] = (
        f"Proyecto: {result.project_id} | {result.enclosure.nombre} | "
        f"Forma {result.separation.forma.value}"
    )
    ws["A2"].font = FONT_NORMAL

    headers = ["Código", "Descripción", "Cantidad", "Unidad", "Categoría", "Observaciones"]
    start = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=start, column=i, value=h)
    _style_header_row(ws, start, len(headers))

    for idx, item in enumerate(result.bom):
        r = start + 1 + idx
        for c, val in enumerate(
            [
                item.codigo,
                item.descripcion,
                item.cantidad,
                item.unidad,
                item.categoria,
                item.observaciones,
            ],
            start=1,
        ):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = FONT_NORMAL
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if idx % 2 == 1:
                cell.fill = FILL_ALT

    ws.freeze_panes = "A5"
    _autosize(ws)


def _write_params(ws, result: ProjectResult) -> None:
    ws["A1"] = "TAN-Telecom — Parámetros Técnicos y Justificación"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:B1")

    rows: list[tuple[str, object]] = [
        ("Identificador de proyecto", result.project_id),
        ("Versión herramienta", result.version),
        ("Generado (UTC)", result.generated_at),
        ("Familia TAN-Telecom", result.separation.familia_tan),
        ("", ""),
        ("— ENTRADAS —", ""),
        ("Potencia instalada [kW]", result.inputs.potencia_kw),
        ("Número de racks / salidas", result.inputs.num_racks),
        ("Redundancia", result.inputs.redundancia.value),
        ("IP mínimo", result.inputs.ip_minimo),
        ("Carga armónica", "Sí" if result.inputs.carga_armonica else "No"),
        ("Requisito sísmico", "Sí" if result.inputs.sismico else "No"),
        ("Tensión [V]", result.inputs.tension_v),
        ("cos φ", result.inputs.cos_phi),
        ("η", result.inputs.rendimiento),
        ("", ""),
        ("— CÁLCULOS ELÉCTRICOS —", ""),
        ("Fórmula In", result.electrical.formula_in),
        ("In base [A]", result.electrical.in_base_a),
        ("Factor de redundancia", result.electrical.factor_redundancia),
        ("In diseño [A]", result.electrical.in_diseno_a),
        ("In selección [A]", result.electrical.in_seleccion_a),
        ("Icc estimada [kA]", result.electrical.icc_estimada_ka),
        ("Icw requerida [kA]", result.electrical.icw_requerida_ka),
        ("RDF", result.electrical.rdf),
        ("Factor de simultaneidad", result.electrical.factor_simultaneidad),
        ("Factor de utilización", result.electrical.factor_utilizacion),
        ("Potencia disipada estimada [W]", result.electrical.potencia_disipada_w),
        ("ΔT estimado [°C]", result.electrical.delta_t_estimado_c),
        ("Límite ΔT [°C]", result.electrical.limite_delta_t_c),
        ("Calentamiento OK", "Sí" if result.electrical.calentamiento_ok else "No"),
        ("", ""),
        ("— ENVOLVENTE —", ""),
        ("Código", result.enclosure.codigo),
        ("Nombre", result.enclosure.nombre),
        ("I max [A]", result.enclosure.corriente_max_a),
        ("Icw / Icc [kA]", f"{result.enclosure.icw_ka} / {result.enclosure.icc_ka}"),
        ("IP seleccionado", result.enclosure.ip_seleccionado),
        ("Dimensiones (ref.)", str(result.enclosure.dimensiones_mm)),
        ("Sísmico catálogo", "Sí" if result.enclosure.resistencia_sismica else "No"),
        ("Alternativas", ", ".join(result.enclosure.alternativas) or "—"),
        ("", ""),
        ("— FORMA DE SEPARACIÓN —", ""),
        ("Forma", result.separation.forma.value),
        ("Criticidad", result.separation.criticidad),
        ("Alternativas", ", ".join(result.separation.alternativas) or "—"),
    ]

    r = 3
    for key, val in rows:
        cell_k = ws.cell(row=r, column=1, value=key)
        cell_v = ws.cell(row=r, column=2, value=val if val != "" else None)
        cell_k.font = FONT_NORMAL
        cell_v.font = FONT_NORMAL
        cell_k.border = THIN
        cell_v.border = THIN
        if key.startswith("—"):
            cell_k.fill = FILL_SUB
            cell_k.font = FONT_HEADER
            cell_v.fill = FILL_SUB
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Justificación envolvente").font = Font(bold=True)
    r += 1
    for line in result.enclosure.justificacion:
        ws.cell(row=r, column=1, value=f"• {line}")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Justificación forma de separación").font = Font(bold=True)
    r += 1
    for line in result.separation.justificacion:
        ws.cell(row=r, column=1, value=f"• {line}")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1

    if result.electrical.notas:
        r += 1
        ws.cell(row=r, column=1, value="Notas de cálculo").font = Font(bold=True)
        r += 1
        for line in result.electrical.notas:
            ws.cell(row=r, column=1, value=f"• {line}")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            r += 1

    if result.warnings:
        r += 1
        ws.cell(row=r, column=1, value="Advertencias").font = Font(bold=True, color="C00000")
        r += 1
        for line in result.warnings:
            cell = ws.cell(row=r, column=1, value=f"• {line}")
            cell.fill = FILL_WARN
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            r += 1

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 70
    ws.freeze_panes = "A3"


def _write_project_link(ws, result: ProjectResult, extras: dict[str, Any]) -> None:
    ws["A1"] = "TAN-Telecom — Proyecto común (Monitor AHF + Banco 24 V)"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:B1")
    observed = extras.get("observed_harmonics") or {}
    bank = extras.get("power_stage") or {}
    rows = [
        ("Proyecto TAN", result.project_id),
        ("Proyecto común", extras.get("shared_project_id") or "—"),
        ("THDi observado [%]", observed.get("thdi")),
        ("Método THDi", observed.get("thdi_method")),
        ("I1 RMS [A]", observed.get("i1_rms")),
        ("H3 RMS [A]", observed.get("h3_rms")),
        ("H5 RMS [A]", observed.get("h5_rms")),
        ("H7 RMS [A]", observed.get("h7_rms")),
        ("In RMS [A]", observed.get("in_rms")),
        ("I compensación [A]", extras.get("compensation_current_a")),
        ("Sugerencia AHF", extras.get("ahf_module_suggestion")),
        ("Nota RIC 04", extras.get("ric04_neutral_note")),
        ("Estado banco 24 V", bank.get("state")),
        ("Falla banco", bank.get("fault")),
        ("Salidas físicas", bank.get("physical_outputs_available")),
    ]
    for idx, (key, val) in enumerate(rows, start=3):
        ws.cell(row=idx, column=1, value=key).font = FONT_NORMAL
        ws.cell(row=idx, column=2, value="" if val is None else val).font = FONT_NORMAL
        ws.cell(row=idx, column=1).border = THIN
        ws.cell(row=idx, column=2).border = THIN
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 80


def generate_excel(
    result: ProjectResult,
    output_path: Path,
    extras: dict[str, Any] | None = None,
) -> Path:
    """
    Genera el libro Excel profesional del proyecto.

    Args:
        result: Resultado de diseño.
        output_path: Ruta del archivo .xlsx.

    Returns:
        Ruta del archivo generado.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Verificaciones de Diseño"
    _write_verifications(ws1, result)

    ws2 = wb.create_sheet("Lista de Materiales (BOM)")
    _write_bom(ws2, result)

    ws3 = wb.create_sheet("Parámetros Técnicos")
    _write_params(ws3, result)

    if extras:
        ws4 = wb.create_sheet("Proyecto comun")
        _write_project_link(ws4, result, extras)

    wb.save(output_path)
    return output_path
